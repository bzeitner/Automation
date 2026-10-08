from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .models import DocumentResult


DOCUMENT_HEADERS = [
    "document_id",
    "status",
    "original_path",
    "original_filename",
    "renamed_path",
    "renamed_filename",
    "sha256",
    "document_type",
    "document_type_label",
    "candidate_document_types",
    "embedded_document_types",
    "document_date",
    "case_number",
    "primary_party",
    "all_parties",
    "extraction_methods",
    "document_type_confidence",
    "overall_confidence",
    "needs_review",
    "review_reasons",
]

FIELD_HEADERS = [
    "document_id",
    "field_key",
    "field_label",
    "value",
    "source_kind",
    "page",
    "evidence",
    "confidence",
    "extraction_method",
]

ERROR_HEADERS = ["document_id", "source_path", "category", "message"]


def _cell_value(value):
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return value


def write_workbook(
    results: list[DocumentResult], errors: list[dict[str, str]], destination: Path
) -> None:
    workbook = Workbook()
    documents_sheet = workbook.active
    documents_sheet.title = "Documents"
    fields_sheet = workbook.create_sheet("Fields")
    errors_sheet = workbook.create_sheet("Errors")
    documents_sheet.append(DOCUMENT_HEADERS)
    fields_sheet.append(FIELD_HEADERS)
    errors_sheet.append(ERROR_HEADERS)

    for index, result in enumerate(results, start=1):
        document_id = f"DOC-{index:06d}"
        documents_sheet.append(
            [
                document_id,
                result.status,
                str(result.source_path),
                result.source_path.name,
                str(result.output_path) if result.output_path else None,
                result.output_path.name if result.output_path else None,
                result.sha256,
                result.document_type,
                result.document_type_label,
                _cell_value(result.candidate_document_types),
                _cell_value(result.embedded_document_types),
                result.document_date,
                result.case_number,
                result.primary_party,
                _cell_value(result.all_parties),
                _cell_value(result.extraction_methods),
                result.document_type_confidence,
                result.overall_confidence,
                result.needs_review,
                _cell_value(result.review_reasons),
            ]
        )
        for field in result.fields:
            fields_sheet.append(
                [
                    document_id,
                    field.key,
                    field.label,
                    _cell_value(field.value),
                    field.source_kind,
                    field.page,
                    field.evidence,
                    field.confidence,
                    field.extraction_method,
                ]
            )

    for error in errors:
        errors_sheet.append([error.get(header) for header in ERROR_HEADERS])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column_cells in sheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max_length + 2, 60)

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
