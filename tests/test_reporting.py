from pathlib import Path

from openpyxl import load_workbook

from document_processor.models import DocumentResult, FieldValue
from document_processor.reporting import write_workbook


def test_workbook_has_expected_sheets(tmp_path):
    result = DocumentResult(
        source_path=Path("source.pdf"),
        sha256="abc",
        document_type="summons",
        document_type_label="Summons",
        fields=[
            FieldValue(
                key="case_number",
                label="Case number",
                value="26-2-02080-18",
                source_kind="explicit",
                confidence=0.99,
                page=1,
                evidence="Case No. 26-2-02080-18",
            )
        ],
    )
    destination = tmp_path / "results.xlsx"
    write_workbook([result], [], destination)
    workbook = load_workbook(destination)

    assert workbook.sheetnames == ["Documents", "Fields", "Errors"]
    assert workbook["Fields"]["D2"].value == "26-2-02080-18"
