from pathlib import Path

from openpyxl import load_workbook
from reportlab.pdfgen import canvas

from document_processor.config import Settings
from document_processor.processor import DocumentProcessor


def _make_pdf(path: Path) -> None:
    pdf = canvas.Canvas(str(path))
    lines = [
        "SUPERIOR COURT OF WASHINGTON FOR KITSAP COUNTY",
        "PINEWOOD MANOR BREMERTON, Plaintiff,",
        "v.",
        "ADRIENNE BARLOW, Defendant.",
        "Case No. 26-2-02080-18",
        "ORDER GRANTING WRIT OF RESTITUTION",
        "DATED October 2, 2026.",
        "This synthetic test document contains no client data. " * 8,
    ]
    y = 750
    for line in lines:
        pdf.drawString(50, y, line)
        y -= 25
    pdf.save()


def _settings() -> Settings:
    return Settings(
        provider="openrouter",
        model=None,
        api_key=None,
        allow_ai=False,
        local_accept_threshold=0.90,
        human_review_threshold=0.85,
    )


def test_processor_preserves_original_and_skips_duplicate(tmp_path, schema):
    input_directory = tmp_path / "intake"
    input_directory.mkdir()
    source = input_directory / "original.pdf"
    _make_pdf(source)
    original_bytes = source.read_bytes()
    output = input_directory / "output"

    processor = DocumentProcessor(
        input_directory=input_directory,
        output_directory=output,
        schema=schema,
        settings=_settings(),
    )
    first_workbook, first_results, first_errors = processor.run()
    second_workbook, second_results, second_errors = processor.run()

    assert source.read_bytes() == original_bytes
    assert first_results[0].status == "review"
    assert "AI was unavailable or disabled" in first_results[0].review_reasons[0]
    assert first_results[0].output_path.exists()
    assert second_results[0].status == "skipped_duplicate"
    assert not first_errors
    assert not second_errors
    assert load_workbook(first_workbook).sheetnames == ["Documents", "Fields", "Errors"]
    assert load_workbook(second_workbook)["Documents"]["B2"].value == "skipped_duplicate"
